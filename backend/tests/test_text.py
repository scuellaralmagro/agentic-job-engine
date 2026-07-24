from pathlib import Path

import pytest

from aje.extraction.text import EmptyExtraction, UnsupportedFileType, extract_text


def _make_pdf(path: Path, body: str) -> None:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, body)
    pdf.output(str(path))


def _make_docx(path: Path, body: str) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph(body)
    doc.save(str(path))


def test_extract_pdf(tmp_path):
    p = tmp_path / "cv.pdf"
    _make_pdf(p, "CurriculumVitae Python Engineer")
    text = extract_text(p)
    assert "Curriculum" in "".join(text.split())


def test_extract_docx(tmp_path):
    p = tmp_path / "cv.docx"
    _make_docx(p, "Senior Backend Developer")
    assert "Senior Backend Developer" in extract_text(p)


def test_unsupported_type(tmp_path):
    p = tmp_path / "cv.txt"
    p.write_text("hi")
    with pytest.raises(UnsupportedFileType):
        extract_text(p)


def test_empty_docx_raises(tmp_path):
    p = tmp_path / "empty.docx"
    _make_docx(p, "")
    with pytest.raises(EmptyExtraction):
        extract_text(p)
